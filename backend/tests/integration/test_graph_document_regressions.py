from io import BytesIO
from types import SimpleNamespace

import pytest
from docx import Document as WordDocument
from fastapi.testclient import TestClient

from app.db import SessionLocal
from app.features.graph import service as graph
from app.features.graph.models import GraphBuild
from app.features.processing.extraction import extract
from app.features.products.models import Revision
from app.main import app
from app.worker import run_once

client = TestClient(app, raise_server_exceptions=False)


def upload_and_extract(name, data):
    product = client.post("/api/products", json={"name": "Graph/document regressions"}).json()
    response = client.post(
        f"/api/products/{product['id']}/documents",
        files={"file": (name, data, "application/octet-stream")},
    )
    assert response.status_code == 201, response.text
    document = response.json()
    # Only extraction and chunking are queued by an upload; no embedding provider is needed.
    run_once("graph-document-regressions")
    run_once("graph-document-regressions")
    return product, client.get(f"/api/documents/{document['id']}").json()


@pytest.mark.parametrize(
    ("name", "data"),
    [
        ("records.csv", b"name,notes\nZulu,Refund requested\n"),
        ("records.json", b'[{"name":"Zulu","notes":"Refund requested"}]'),
        (
            "records.ttl",
            b'@prefix ex: <https://records.example/> . ex:zulu a ex:Person; ex:knows ex:ana . ex:ana a ex:Person .',
        ),
    ],
)
def test_structured_build_reports_missing_property_mapping_and_can_recover(name, data):
    product, document = upload_and_extract(name, data)
    pid = product["id"]
    definitions = client.get(f"/api/products/{pid}/ontology").json()
    complete_mapping = definitions["mapping"]
    response = client.put(
        f"/api/products/{pid}/ontology/mapping",
        json={"classes": complete_mapping["classes"], "properties": []},
    )
    assert response.status_code == 200, response.text
    response = client.post(f"/api/products/{pid}/graph/build")
    assert response.status_code == 409, response.text
    detail = response.json()["detail"]
    assert detail["code"] == "unmapped_properties"
    assert "representation" in detail["message"].lower()
    assert "advanced" in detail["message"].lower()
    assert client.get(f"/api/products/{pid}").json()["draft"]["graph_build_id"] is None
    assert client.get(f"/api/documents/{document['id']}/original").content == data

    response = client.put(
        f"/api/products/{pid}/ontology/mapping",
        json={"classes": complete_mapping["classes"], "properties": complete_mapping["properties"]},
    )
    assert response.status_code == 200, response.text
    response = client.post(f"/api/products/{pid}/graph/build")
    assert response.status_code == 200, response.text
    nodes = client.get(f"/api/products/{pid}/entities").json()["items"]
    assert all(node["evidence"][0]["document_id"] == document["id"] for node in nodes)


def colliding_graph():
    names = {1: "Zulu", 2: "Yankee", 10: "Alpha", 20: "Aardvark"}
    data = ("name\n" + "\n".join(names.get(i, f"Record {i}") for i in range(1, 21))).encode()
    product, document = upload_and_extract("colliding.csv", data)
    built = client.post(f"/api/products/{product['id']}/graph/build")
    assert built.status_code == 200, built.text
    nodes = client.get(f"/api/products/{product['id']}/entities").json()["items"]
    root = next(node for node in nodes if node["label"] == "Zulu")
    neighbor = next(node for node in nodes if node["label"] == "Yankee")
    with SessionLocal() as session:
        build = session.get(GraphBuild, built.json()["id"])
        build.relationships = [
            {
                "source": root["id"],
                "target": neighbor["id"],
                "type": "RELATED_TO",
                "evidence": root["evidence"],
            }
        ]
        graph.adapter.build(build)
        session.commit()
    return product, document, built.json(), root, neighbor


def test_neighborhood_uses_exact_root_and_neighbor_ids_without_changing_search():
    product, _, _, root, neighbor = colliding_graph()
    pid = product["id"]
    search = client.get(f"/api/products/{pid}/entities", params={"q": root["id"]}).json()["items"]
    assert {"Alpha", "Zulu"} <= {node["label"] for node in search}
    assert len(search) > 1
    response = client.get(f"/api/products/{pid}/entities/{root['id']}/neighbors?limit=2")
    assert response.status_code == 200, response.text
    neighborhood = response.json()
    assert [node["id"] for node in neighborhood["nodes"]] == [root["id"], neighbor["id"]]
    assert neighborhood["nodes"][0]["evidence"] == root["evidence"]
    assert neighborhood["nodes"][1]["evidence"] == neighbor["evidence"]
    assert neighborhood["relationships"][0]["source"] == root["id"]
    assert neighborhood["relationships"][0]["target"] == neighbor["id"]
    limited = client.get(f"/api/products/{pid}/entities/{root['id']}/neighbors?limit=1").json()
    assert [node["id"] for node in limited["nodes"]] == [root["id"]]
    assert limited["relationships"] == []


def test_hybrid_graph_context_keeps_the_fact_from_its_matching_excerpt():
    product, _, _, root, neighbor = colliding_graph()
    with SessionLocal() as session:
        revision = session.get(Revision, product["draft"]["id"])
        contexts = graph.graph_context(session, revision, root["evidence"][0]["chunk_id"], 2)
    assert len(contexts) == 1
    assert [node["id"] for node in contexts[0]["nodes"]] == [root["id"], neighbor["id"]]


def test_neighborhood_returns_404_for_a_missing_exact_id_despite_substring_matches():
    product, _, _, root, _ = colliding_graph()
    # The dataset prefix matches all records, but is not itself an entity ID.
    missing_id = root["id"].rsplit("row-", 1)[0]
    response = client.get(f"/api/products/{product['id']}/entities/{missing_id}/neighbors")
    assert response.status_code == 404, response.text


def word_bytes(document):
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_table_only_docx_extracts_readable_text_and_cited_chunks():
    word = WordDocument()
    table = word.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Complaint C-1042"
    table.cell(0, 1).text = "Refund requested"
    data = word_bytes(word)
    _, document = upload_and_extract("table-only.docx", data)
    assert document["extracted_text"] is not None, document["jobs"]
    assert document["extracted_text"].splitlines() == ["Complaint C-1042", "Refund requested"]
    assert document["chunks"][0]["text"] == document["extracted_text"]
    assert document["chunks"][0]["start"] == 0
    assert document["chunks"][0]["end"] == len(document["extracted_text"])
    assert document["processing_version"] == "extract-docx-v2/chunk-v1"
    assert client.get(f"/api/documents/{document['id']}/original").content == data


def test_docx_preserves_paragraph_table_and_nested_table_reading_order():
    word = WordDocument()
    word.add_paragraph("Before table")
    table = word.add_table(rows=1, cols=2)
    cell = table.cell(0, 0)
    cell.text = "Cell before nested table"
    cell.add_table(rows=1, cols=1).cell(0, 0).text = "Nested value"
    cell.add_paragraph("Cell after nested table")
    table.cell(0, 1).text = "Second cell"
    word.add_paragraph("After table")
    _, document = upload_and_extract("mixed.docx", word_bytes(word))
    assert document["extracted_text"] is not None, document["jobs"]
    assert [line for line in document["extracted_text"].splitlines() if line] == [
        "Before table",
        "Cell before nested table",
        "Nested value",
        "Cell after nested table",
        "Second cell",
        "After table",
    ]
    assert document["chunks"][0]["text"] == document["extracted_text"]


def test_previously_extracted_docx_preserves_its_text_hash_and_processing_version():
    document = SimpleNamespace(
        name="prepared.docx",
        extracted_key="previous-text",
        extracted_text="Previously captured paragraph evidence",
        extracted_sha256="previous-hash",
        processing_version="extract-v1/chunk-v1",
    )
    extract(None, document)
    assert document.extracted_key == "previous-text"
    assert document.extracted_text == "Previously captured paragraph evidence"
    assert document.extracted_sha256 == "previous-hash"
    assert document.processing_version == "extract-v1/chunk-v1"
